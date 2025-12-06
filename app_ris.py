# import streamlit as st
# import pandas as pd
# import re
# from io import BytesIO
# from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
# from openpyxl.utils import get_column_letter


# st.set_page_config(page_title="RIS Analysis Tool", page_icon="📊", layout="wide")

# st.markdown(
#     """
#     <style>
#     /* Center main content and limit width */
#     .main > div {
#         max-width: 1200px;
#         margin-left: auto;
#         margin-right: auto;
#     }
#     .app-card {
#         background-color: #111827;
#         border-radius: 12px;
#         padding: 1.1rem 1.3rem;
#         border: 1px solid #1f2937;
#         box-shadow: 0 2px 6px rgba(15,23,42,0.4);
#         margin-bottom: 1.25rem;
#     }
#     .section-title {
#         font-size: 1.2rem;
#         font-weight: 700;
#         margin-bottom: 0.35rem;
#     }
#     .section-subtitle {
#         font-size: 0.9rem;
#         color: #9ca3af;
#         margin-bottom: 0.6rem;
#     }
#     .step-pill {
#         display: inline-block;
#         padding: 0.10rem 0.6rem;
#         border-radius: 999px;
#         background: #1f2937;
#         color: #93c5fd;
#         font-size: 0.78rem;
#         font-weight: 600;
#         margin-bottom: 0.3rem;
#     }
#     </style>
#     """,
#     unsafe_allow_html=True,
# )

# st.title("📊 RIS Analysis Tool")
# st.markdown(
#     "Upload **Original.xlsx**, **FC Stat.xlsx** and **PM.xlsx (Purchase Master)**. "
#     "This app will use `Original.sku` matched to PM Column C → Brand (Column G)."
# )

# # --------------------------
# # Helpers
# # --------------------------
# def normalize(s):
#     return s.astype(str).str.strip().str.upper().replace({"": pd.NA})

# def sku_normalize_raw(s):
#     """Normalize SKU: trim, uppercase, remove WO- prefix, remove non-alphanumeric"""
#     if pd.isna(s):
#         return ""
#     s = str(s).strip().upper()
#     s = re.sub(r'^WO[-_]?','', s)           
#     s = re.sub(r'[^A-Z0-9]', '', s)        
#     return s

# def normalize_text(s):
#     if pd.isna(s):
#         return ""
#     s = str(s).replace("\xa0", " ")
#     s = s.lower().strip()
#     s = re.sub(r"[^\w\s]", "", s)
#     s = re.sub(r"\s+", " ", s).strip()
#     return s.replace(" ", "")

# def safe_correct(raw, canon_map):
#     norm = normalize_text(raw)
#     return canon_map.get(norm, raw)

# def build_cluster_display_with_pct(cluster_df):
#     if cluster_df is None or cluster_df.empty:
#         return pd.DataFrame()
#     cols = [c for c in ["Non RIS", "RIS", "Grand Total"] if c in cluster_df.columns]
#     rows = []
#     idx = []
#     clusters = []
#     for t in cluster_df.index:
#         if t[0] not in clusters:
#             clusters.append(t[0])
#     for cl in clusters:
#         try:
#             sub = cluster_df.loc[cl]
#         except KeyError:
#             continue
#         if isinstance(sub, pd.Series):
#             sub_df = pd.DataFrame([sub]); sub_df.index = [sub.name]
#         else:
#             sub_df = sub.copy()
#         rows.append([pd.NA]* (len(cols)+2))
#         idx.append((cl, "", ""))
#         for brand in sub_df.index:
#             nonris = float(sub_df.at[brand, "Non RIS"]) if "Non RIS" in sub_df.columns else 0.0
#             ris = float(sub_df.at[brand, "RIS"]) if "RIS" in sub_df.columns else 0.0
#             grand = nonris + ris
#             nonris_pct = (nonris / grand) if grand != 0 else 0.0
#             ris_pct = (ris / grand) if grand != 0 else 0.0
#             rows.append([nonris, ris, grand, nonris_pct, ris_pct])
#             idx.append((cl, brand, ""))
#         sc_total = sub_df[cols].sum(axis=0)
#         sc_non = float(sc_total.get("Non RIS", 0.0))
#         sc_ris = float(sc_total.get("RIS", 0.0))
#         sc_grand = sc_non + sc_ris
#         sc_non_pct = (sc_non / sc_grand) if sc_grand != 0 else 0.0
#         sc_ris_pct = (sc_ris / sc_grand) if sc_grand != 0 else 0.0
#         rows.append([sc_non, sc_ris, sc_grand, sc_non_pct, sc_ris_pct])
#         idx.append((cl, f"{cl} Cluster Total", ""))
#     try:
#         grand = cluster_df[[c for c in ["Non RIS", "RIS"] if c in cluster_df.columns]].sum(axis=0)
#         g_non = float(grand.get("Non RIS", 0.0))
#         g_ris = float(grand.get("RIS", 0.0))
#         g_grand = g_non + g_ris
#         g_non_pct = (g_non / g_grand) if g_grand != 0 else 0.0
#         g_ris_pct = (g_ris / g_grand) if g_grand != 0 else 0.0
#         rows.append([g_non, g_ris, g_grand, g_non_pct, g_ris_pct])
#         idx.append(("Grand Total", "Grand Total", ""))
#     except Exception:
#         pass
#     display_cols = cols + ["Non RIS %", "RIS %"]
#     df = pd.DataFrame(
#         rows,
#         columns=display_cols,
#         index=pd.MultiIndex.from_tuples(idx, names=["State Cluster", "Brand", ""])
#     )
#     return df

# def style_cluster_sheet_openpyxl(writer, sheet_name):
#     ws = writer.sheets[sheet_name]
#     thin = Side(border_style="thin", color="000000")
#     border = Border(left=thin, right=thin, top=thin, bottom=thin)
#     max_row = ws.max_row; max_col = ws.max_column
#     for r in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
#         for cell in r:
#             cell.border = border
#             if isinstance(cell.value, (int, float)):
#                 cell.alignment = Alignment(horizontal="center", vertical="center")
#             else:
#                 cell.alignment = Alignment(horizontal="left", vertical="center")
#     for row in range(2, max_row+1):
#         brand_val = ws.cell(row=row, column=2).value
#         try:
#             if brand_val is None or str(brand_val).strip() == "":
#                 for col in range(1, max_col+1):
#                     cell = ws.cell(row=row, column=col)
#                     cell.font = Font(bold=True)
#                     cell.fill = PatternFill(start_color="B7CDE6", end_color="B7CDE6", fill_type="solid")
#                 a_cell = ws.cell(row=row, column=1); a_cell.alignment = Alignment(horizontal="left", indent=0)
#             elif "Cluster Total" in str(brand_val):
#                 for col in range(1, max_col+1):
#                     cell = ws.cell(row=row, column=col); cell.font = Font(bold=True); cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
#             elif "Grand Total" in str(brand_val):
#                 for col in range(1, max_col+1):
#                     cell = ws.cell(row=row, column=col); cell.font = Font(bold=True); cell.fill = PatternFill(start_color="9FB3C8", end_color="9FB3C8", fill_type="solid")
#         except Exception:
#             pass
#     for row in range(2, max_row+1):
#         bcell = ws.cell(row=row, column=2)
#         if bcell.value and "Cluster Total" not in str(bcell.value) and "Grand Total" not in str(bcell.value):
#             bcell.alignment = Alignment(horizontal="left", indent=2)
#     for col in range(1, max_col+1):
#         col_letter = get_column_letter(col)
#         ws.column_dimensions[col_letter].width = 18

# def apply_number_formats(ws, percent_cols_idx, number_cols_idx, header_row=1):
#     max_row = ws.max_row
#     for r in range(header_row + 1, max_row + 1):
#         for c in percent_cols_idx:
#             try:
#                 cell = ws.cell(row=r, column=c); cell.number_format = '0.00%'
#             except Exception:
#                 pass
#         for c in number_cols_idx:
#             try:
#                 cell = ws.cell(row=r, column=c); cell.number_format = '#,##0'
#             except Exception:
#                 pass


# st.markdown('<div class="app-card">', unsafe_allow_html=True)
# st.markdown('<span class="step-pill">STEP 1</span>', unsafe_allow_html=True)
# st.markdown('<div class="section-title">Upload Files</div>', unsafe_allow_html=True)
# st.markdown(
#     '<div class="section-subtitle">'
#     'Upload the required Excel files to start the RIS analysis.'
#     '</div>',
#     unsafe_allow_html=True,
# )

# c1, c2, c3 = st.columns(3)

# with c1:
#     st.markdown("**Original.xlsx**")
#     original_file = st.file_uploader("Upload Original.xlsx", type=["xlsx"], key="orig_up")

# with c2:
#     st.markdown("**FC Stat.xlsx**")
#     fc_file = st.file_uploader("Upload FC Stat.xlsx", type=["xlsx"], key="fc_up")

# with c3:
#     st.markdown("**PM.xlsx (Purchase Master)**")
#     pm_file = st.file_uploader("Upload PM.xlsx", type=["xlsx"], key="pm_up")

# st.markdown("</div>", unsafe_allow_html=True)

# if original_file and fc_file and pm_file:
#     try:
#         with st.spinner("Loading data..."):
#             original_file.seek(0); Working = pd.read_excel(original_file)
#             fc_file.seek(0); FC = pd.read_excel(fc_file)
#             pm_file.seek(0)
#             try:
#                 PM = pd.read_excel(pm_file, usecols="C,G", sheet_name=0)
#             except Exception:
#                 PM = pd.read_excel(pm_file)

#         st.success("Files loaded ✅")
#         st.write(f"Original rows: {len(Working)} — FC rows: {len(FC)} — PM rows: {len(PM)}")

#         if st.button("🔄 Process Data"):
#             with st.spinner("Processing..."):
#                 # -- basic checks
#                 key_col_name = "fulfillment-center-id"
#                 if key_col_name not in Working.columns:
#                     st.error(f"Missing column: {key_col_name}"); st.stop()
#                 if "quantity-shipped" not in Working.columns:
#                     st.error("Missing 'quantity-shipped' column in Original file."); st.stop()

#                 # ensure numeric measure
#                 Working["quantity-shipped"] = pd.to_numeric(Working["quantity-shipped"], errors="coerce").fillna(0)

#                 # Merge FC info (unique keys only)
#                 FC_lookup = FC[["FC", "State", "Cluster"]].copy()
#                 FC_lookup.columns = ["lookup_key", "fulfillment_state", "cluster"]
#                 Working["_join"] = normalize(Working[key_col_name])
#                 FC_lookup["_join"] = normalize(FC_lookup["lookup_key"])
#                 FC_lookup = FC_lookup.drop_duplicates(subset=["_join"], keep="first")
#                 Working = Working.merge(
#                     FC_lookup[["_join", "fulfillment_state", "cluster"]],
#                     how="left",
#                     left_on="_join",
#                     right_on="_join"
#                 )
#                 Working.drop(columns=["_join"], inplace=True)

#                 # Normalize ship-state
#                 fc_states = FC["State"].dropna().astype(str).str.strip().tolist()
#                 canon_map = {normalize_text(s): s.strip() for s in fc_states if normalize_text(s) != ""}
#                 if "ship-state_original" not in Working.columns:
#                     Working["ship-state_original"] = Working.get("ship-state", pd.NA)
#                 Working["ship-state"] = Working["ship-state_original"].apply(lambda x: safe_correct(x, canon_map))

#                 # Sheet2 VLOOKUP (A:C) using fulfillment-center-id if present
#                 try:
#                     original_file.seek(0); xl = pd.ExcelFile(original_file)
#                     sheet2 = xl.parse(xl.sheet_names[1]) if len(xl.sheet_names) > 1 else None
#                 except Exception:
#                     sheet2 = None

#                 if sheet2 is not None and sheet2.shape[1] >= 3:
#                     map_df = sheet2.iloc[:, [0, 2]].copy(); map_df.columns = ["key_raw", "state_cluster_raw"]
#                     map_df["_k"] = map_df["key_raw"].astype(str).apply(lambda x: str(x).strip().upper())
#                     mapping = dict(zip(map_df["_k"], map_df["state_cluster_raw"]))
#                     Working["State Cluster"] = Working[key_col_name].astype(str).apply(
#                         lambda x: mapping.get(str(x).strip().upper(), pd.NA)
#                     )
#                 else:
#                     Working["State Cluster"] = pd.NA

#                 # RIS status
#                 Working["RIS Status"] = Working.apply(
#                     lambda row: "RIS"
#                     if str(row.get("ship-state_original", "")).strip().replace(" ", "")
#                        == str(row.get("fulfillment_state", "")).strip().replace(" ", "")
#                     else "Non RIS",
#                     axis=1
#                 )

#                 # --------------------------
#                 # PM -> Brand mapping (Original.sku  <-> PM Column C -> Brand Column G)
#                 # --------------------------
#                 PM_lookup = pd.DataFrame()
#                 try:
#                     if PM.shape[1] >= 2:
#                         PM_lookup = PM.iloc[:, [0, 1]].copy()
#                         PM_lookup.columns = ["lookup_key", "Brand_pm"]
#                     else:
#                         PM_lookup = pd.DataFrame()
#                 except Exception:
#                     PM_lookup = pd.DataFrame()

#                 if not PM_lookup.empty:
#                     PM_lookup["lookup_key_norm"] = PM_lookup["lookup_key"].apply(sku_normalize_raw)
#                     PM_lookup["Brand_pm"] = PM_lookup["Brand_pm"].astype(str).str.strip()
#                     PM_lookup = PM_lookup.drop_duplicates(subset=["lookup_key_norm"]).reset_index(drop=True)

#                 # Normalize Working.sku similarly
#                 if "sku" in Working.columns:
#                     Working["_sku_norm"] = Working["sku"].apply(sku_normalize_raw)
#                 else:
#                     Working["_sku_norm"] = ""

#                 # Ensure Brand present and preserve existing values: fill only missing
#                 if "Brand" not in Working.columns:
#                     Working["Brand"] = ""
#                 Working["Brand"] = Working["Brand"].fillna("").astype(str)

#                 filled_count = 0
#                 if not PM_lookup.empty:
#                     pm_map = dict(zip(PM_lookup["lookup_key_norm"], PM_lookup["Brand_pm"]))
#                     pm_keys_set = set(pm_map.keys())
#                     mask_unknown = Working["Brand"].str.strip().str.upper().isin(["", "NAN", "NONE", "UNKNOWN"])
#                     match_mask = Working["_sku_norm"].isin(pm_keys_set) & mask_unknown
#                     if match_mask.any():
#                         for idx in Working[match_mask].index:
#                             k = Working.at[idx, "_sku_norm"]
#                             val = pm_map.get(k, "")
#                             if val:
#                                 Working.at[idx, "Brand"] = val
#                                 filled_count += 1

#                 Working["Brand"] = Working["Brand"].replace("", "Unknown")

#                 st.markdown("### PM→Brand diagnostics")
#                 st.write("PM lookup rows (after normalize & dedup):", len(PM_lookup))
#                 st.write("Filled Brand rows from PM:", int(filled_count))
#                 st.write("Distinct Brands after fill:", int(Working["Brand"].nunique()))
#                 st.write(
#                     "Rows remaining with Brand == 'Unknown':",
#                     int((Working["Brand"].astype(str).str.strip().str.upper() == "UNKNOWN").sum())
#                 )

#                 if "_sku_norm" in Working.columns:
#                     try:
#                         Working.drop(columns=["_sku_norm"], inplace=True)
#                     except Exception:
#                         pass

#                 # --------------------------
#                 # Pivots
#                 # --------------------------
#                 detailed_pivot = pd.pivot_table(
#                     Working,
#                     values="quantity-shipped",
#                     index=["Brand", "fulfillment_state", "ship-state"],
#                     columns="RIS Status",
#                     aggfunc="sum",
#                     fill_value=0,
#                 )

#                 if not detailed_pivot.empty:
#                     if "Non RIS" not in detailed_pivot.columns: detailed_pivot["Non RIS"] = 0
#                     if "RIS" not in detailed_pivot.columns: detailed_pivot["RIS"] = 0
#                     detailed_pivot["Grand Total"] = detailed_pivot["Non RIS"] + detailed_pivot["RIS"]

#                     brand_totals = detailed_pivot.groupby(level=0).sum()
#                     brand_totals.index = [(brand, f"{brand} Total", "") for brand in brand_totals.index]

#                     state_totals = detailed_pivot.groupby(level=[0, 1]).sum()
#                     state_totals.index = [(brand, f"{state} Total", "") for brand, state in state_totals.index]

#                     detailed_pivot_with_totals = pd.concat(
#                         [detailed_pivot, state_totals, brand_totals]
#                     ).sort_index()

#                     overall = detailed_pivot[["Non RIS", "RIS", "Grand Total"]].sum()
#                     overall_idx = ("Grand Total", "", "")
#                     overall_df = pd.DataFrame([overall.values], index=[overall_idx], columns=overall.index)
#                     detailed_pivot_with_totals = pd.concat([detailed_pivot_with_totals, overall_df])
#                 else:
#                     detailed_pivot_with_totals = detailed_pivot.copy()

#                 brand_summary = pd.pivot_table(
#                     Working,
#                     values="quantity-shipped",
#                     index=["Brand"],
#                     columns="RIS Status",
#                     aggfunc="sum",
#                     fill_value=0,
#                     margins=False,
#                 )
#                 if "Non RIS" not in brand_summary.columns: brand_summary["Non RIS"] = 0
#                 if "RIS" not in brand_summary.columns: brand_summary["RIS"] = 0
#                 brand_summary["Grand Total"] = brand_summary["Non RIS"] + brand_summary["RIS"]
#                 brand_summary["Non RIS%"] = (brand_summary["Non RIS"] / brand_summary["Grand Total"]).fillna(0)
#                 brand_summary["RIS%"] = (brand_summary["RIS"] / brand_summary["Grand Total"]).fillna(0)

#                 overall_brand = brand_summary[["Non RIS", "RIS", "Grand Total"]].sum()
#                 overall_brand["Non RIS%"] = (
#                     overall_brand["Non RIS"] / overall_brand["Grand Total"]
#                     if overall_brand["Grand Total"] != 0 else 0
#                 )
#                 overall_brand["RIS%"] = (
#                     overall_brand["RIS"] / overall_brand["Grand Total"]
#                     if overall_brand["Grand Total"] != 0 else 0
#                 )
#                 overall_brand = overall_brand.rename("Grand Total")
#                 overall_brand_df = overall_brand.to_frame().T
#                 brand_summary = pd.concat([brand_summary, overall_brand_df])

#                 inventory_brand_summary = pd.pivot_table(
#                     Working,
#                     values="quantity-shipped",
#                     index=["Brand"],
#                     columns="RIS_by_Table",
#                     aggfunc="sum",
#                     fill_value=0,
#                     margins=True,
#                     margins_name="Grand Total",
#                 ) if "RIS_by_Table" in Working.columns else pd.DataFrame()

#                 if not inventory_brand_summary.empty:
#                     if "Non RIS" not in inventory_brand_summary.columns: inventory_brand_summary["Non RIS"] = 0
#                     if "RIS" not in inventory_brand_summary.columns: inventory_brand_summary["RIS"] = 0
#                     inventory_brand_summary["Grand Total"] = inventory_brand_summary.get(
#                         "Grand Total",
#                         inventory_brand_summary["Non RIS"] + inventory_brand_summary["RIS"],
#                     )
#                     inventory_brand_summary["Non RIS%"] = (
#                         inventory_brand_summary["Non RIS"] / inventory_brand_summary["Grand Total"]
#                     ).fillna(0)
#                     inventory_brand_summary["RIS%"] = (
#                         inventory_brand_summary["RIS"] / inventory_brand_summary["Grand Total"]
#                     ).fillna(0)

#                 cluster_pivot = pd.pivot_table(
#                     Working,
#                     values="quantity-shipped",
#                     index=["State Cluster", "Brand"],
#                     columns="RIS Status",
#                     aggfunc="sum",
#                     fill_value=0,
#                 )
#                 if not cluster_pivot.empty:
#                     if "Non RIS" not in cluster_pivot.columns: cluster_pivot["Non RIS"] = 0
#                     if "RIS" not in cluster_pivot.columns: cluster_pivot["RIS"] = 0
#                     cluster_pivot["Grand Total"] = cluster_pivot["Non RIS"] + cluster_pivot["RIS"]

#                 display_cluster_df = build_cluster_display_with_pct(cluster_pivot)

#             st.success("✅ Processing complete!")

#             # --------------------------
#             # UI: Tabs & displays
#             # --------------------------
#             tab1, tab2 = st.tabs(["📊 State-Based RIS Analysis", "🏭 Inventory Placement"])
#             with tab1:
#                 st.subheader("Brand-Level Summary (with %)")
#                 st.dataframe(
#                     brand_summary.style.format({
#                         "Non RIS": "{:,.0f}",
#                         "RIS": "{:,.0f}",
#                         "Grand Total": "{:,.0f}",
#                         "Non RIS%": "{:.2%}",
#                         "RIS%": "{:.2%}",
#                     }),
#                     use_container_width=True,
#                 )

#                 st.subheader("Detailed (Brand → FC State → Ship State)")
#                 st.dataframe(
#                     detailed_pivot_with_totals.style.format({
#                         "Non RIS": "{:,.0f}",
#                         "RIS": "{:,.0f}",
#                         "Grand Total": "{:,.0f}",
#                     }),
#                     use_container_width=True,
#                     height=400,
#                 )

#                 st.subheader("Cluster Analysis — Excel-like structure")
#                 if display_cluster_df.empty:
#                     st.info("No cluster data available (Sheet2 missing or no matches).")
#                 else:
#                     def style_rows(s):
#                         label = s.name[1]
#                         if label == "":
#                             return ["background-color: #B7CDE6; font-weight: bold; text-align:left" for _ in s]
#                         elif "Cluster Total" in str(label):
#                             return ["background-color: #E0E0E0; font-weight: bold;" for _ in s]
#                         elif "Grand Total" in str(label):
#                             return ["background-color: #9FB3C8; font-weight: bold;" for _ in s]
#                         else:
#                             return ["" for _ in s]

#                     st.dataframe(
#                         display_cluster_df.style.apply(style_rows, axis=1).format({
#                             "Non RIS": "{:,.0f}",
#                             "RIS": "{:,.0f}",
#                             "Grand Total": "{:,.0f}",
#                             "Non RIS %": "{:.2%}",
#                             "RIS %": "{:.2%}",
#                         }),
#                         use_container_width=True,
#                         height=650,
#                     )

#             with tab2:
#                 st.header("Inventory Placement (Brand summary)")
#                 if not inventory_brand_summary.empty:
#                     st.dataframe(
#                         inventory_brand_summary.style.format({
#                             "Non RIS": "{:,.0f}",
#                             "RIS": "{:,.0f}",
#                             "Grand Total": "{:,.0f}",
#                             "Non RIS%": "{:.2%}",
#                             "RIS%": "{:.2%}",
#                         }),
#                         use_container_width=True,
#                     )
#                 else:
#                     st.info("No Inventory Placement summary (RIS_by_Table missing).")

#             # --------------------------
#             # Downloads
#             # --------------------------
#             st.subheader("📥 Download Report (Excel)")
#             col1, col2 = st.columns(2)
#             with col1:
#                 out = BytesIO()
#                 with pd.ExcelWriter(out, engine="openpyxl") as writer:
#                     brand_summary.to_excel(writer, sheet_name="Brand Summary")
#                     detailed_pivot_with_totals.to_excel(writer, sheet_name="Detailed Analysis")
#                     if not inventory_brand_summary.empty:
#                         inventory_brand_summary.to_excel(writer, sheet_name="Inventory Brand Summary")
#                     if not display_cluster_df.empty:
#                         display_cluster_df.to_excel(writer, sheet_name="StateCluster Pivot")

#                     for sheet_name in writer.sheets:
#                         ws = writer.sheets[sheet_name]
#                         for cell in ws[1]:
#                             cell.font = Font(bold=True)

#                         header_map = {
#                             col_idx + 1: ws.cell(row=1, column=col_idx + 1).value
#                             for col_idx in range(ws.max_column)
#                         }

#                         percent_cols = [
#                             i for i, v in header_map.items()
#                             if v in ["Non RIS%", "RIS%", "Non RIS %", "RIS %"]
#                         ]

#                         number_cols = [
#                             i for i, v in header_map.items()
#                             if isinstance(v, str)
#                             and any(k in v for k in ["Non RIS", "RIS", "Grand Total"])
#                             and "%" not in str(v)
#                         ]

#                         apply_number_formats(
#                             ws,
#                             percent_cols_idx=percent_cols,
#                             number_cols_idx=number_cols,
#                             header_row=1,
#                         )

#                         if sheet_name == "StateCluster Pivot":
#                             style_cluster_sheet_openpyxl(writer, sheet_name)

#                         for col in ws.columns:
#                             max_len = 0
#                             col_letter = get_column_letter(col[0].column)
#                             for cell in col:
#                                 try:
#                                     if cell.value:
#                                         max_len = max(max_len, len(str(cell.value)))
#                                 except Exception:
#                                     pass
#                             ws.column_dimensions[col_letter].width = max(12, max_len + 2)

#                 out.seek(0)
#                 st.download_button(
#                     "📦 Download Excel report",
#                     data=out,
#                     file_name="ris_cluster_report.xlsx",
#                     mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#                 )
#             with col2:
#                 st.download_button(
#                     "📄 Download Processed Data (CSV)",
#                     data=Working.to_csv(index=False).encode("utf-8"),
#                     file_name="processed_data.csv",
#                     mime="text/csv",
#                 )

#     except Exception as e:
#         st.error("Error: " + str(e))
#         st.exception(e)

# else:
#     st.info("Please upload **Original.xlsx**, **FC Stat.xlsx** and **PM.xlsx** above to proceed.")
#     st.markdown(
#         """
#         - `Original.xlsx` must include **Sheet2** where Column A = `fulfillment-center-id`
#           and Column C = **State Cluster** (used for VLOOKUP style).
#         - `fulfillment-center-id` column is required in the Original sheet.
#         - `quantity-shipped` is used as measure.
#         """
#     )

import streamlit as st
import pandas as pd
import re
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


st.set_page_config(page_title="RIS Analysis Tool", page_icon="📊", layout="wide")

st.markdown(
    """
    <style>
    /* Center main content and limit width */
    .main > div {
        max-width: 1200px;
        margin-left: auto;
        margin-right: auto;
    }
    .app-card {
        background-color: #111827;
        border-radius: 12px;
        padding: 1.1rem 1.3rem;
        border: 1px solid #1f2937;
        box-shadow: 0 2px 6px rgba(15,23,42,0.4);
        margin-bottom: 1.25rem;
    }
    .section-title {
        font-size: 1.2rem;
        font-weight: 700;
        margin-bottom: 0.35rem;
    }
    .section-subtitle {
        font-size: 0.9rem;
        color: #9ca3af;
        margin-bottom: 0.6rem;
    }
    .step-pill {
        display: inline-block;
        padding: 0.10rem 0.6rem;
        border-radius: 999px;
        background: #1f2937;
        color: #93c5fd;
        font-size: 0.78rem;
        font-weight: 600;
        margin-bottom: 0.3rem;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("📊 RIS Analysis Tool")
st.markdown(
    "Upload **Original.xlsx**, **FC Stat.xlsx** and **PM.xlsx (Purchase Master)**. "
    "This app will use `Original.sku` matched to PM Column C → Brand (Column G)."
)

# --------------------------
# Helpers
# --------------------------
def normalize(s):
    return s.astype(str).str.strip().str.upper().replace({"": pd.NA})

def sku_normalize_raw(s):
    """Normalize SKU: trim, uppercase, remove WO- prefix, remove non-alphanumeric"""
    if pd.isna(s):
        return ""
    s = str(s).strip().upper()
    s = re.sub(r'^WO[-_]?','', s)           
    s = re.sub(r'[^A-Z0-9]', '', s)        
    return s

def normalize_text(s):
    if pd.isna(s):
        return ""
    s = str(s).replace("\xa0", " ")
    s = s.lower().strip()
    s = re.sub(r"[^\w\s]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s.replace(" ", "")

def safe_correct(raw, canon_map):
    norm = normalize_text(raw)
    return canon_map.get(norm, raw)

def build_cluster_display_with_pct(cluster_df):
    if cluster_df is None or cluster_df.empty:
        return pd.DataFrame()
    cols = [c for c in ["Non RIS", "RIS", "Grand Total"] if c in cluster_df.columns]
    rows = []
    idx = []
    clusters = []
    for t in cluster_df.index:
        if t[0] not in clusters:
            clusters.append(t[0])
    for cl in clusters:
        try:
            sub = cluster_df.loc[cl]
        except KeyError:
            continue
        if isinstance(sub, pd.Series):
            sub_df = pd.DataFrame([sub]); sub_df.index = [sub.name]
        else:
            sub_df = sub.copy()
        rows.append([pd.NA]* (len(cols)+2))
        idx.append((cl, "", ""))
        for brand in sub_df.index:
            nonris = float(sub_df.at[brand, "Non RIS"]) if "Non RIS" in sub_df.columns else 0.0
            ris = float(sub_df.at[brand, "RIS"]) if "RIS" in sub_df.columns else 0.0
            grand = nonris + ris
            nonris_pct = (nonris / grand) if grand != 0 else 0.0
            ris_pct = (ris / grand) if grand != 0 else 0.0
            rows.append([nonris, ris, grand, nonris_pct, ris_pct])
            idx.append((cl, brand, ""))
        sc_total = sub_df[cols].sum(axis=0)
        sc_non = float(sc_total.get("Non RIS", 0.0))
        sc_ris = float(sc_total.get("RIS", 0.0))
        sc_grand = sc_non + sc_ris
        sc_non_pct = (sc_non / sc_grand) if sc_grand != 0 else 0.0
        sc_ris_pct = (sc_ris / sc_grand) if sc_grand != 0 else 0.0
        rows.append([sc_non, sc_ris, sc_grand, sc_non_pct, sc_ris_pct])
        idx.append((cl, f"{cl} Cluster Total", ""))
    try:
        grand = cluster_df[[c for c in ["Non RIS", "RIS"] if c in cluster_df.columns]].sum(axis=0)
        g_non = float(grand.get("Non RIS", 0.0))
        g_ris = float(grand.get("RIS", 0.0))
        g_grand = g_non + g_ris
        g_non_pct = (g_non / g_grand) if g_grand != 0 else 0.0
        g_ris_pct = (g_ris / g_grand) if g_grand != 0 else 0.0
        rows.append([g_non, g_ris, g_grand, g_non_pct, g_ris_pct])
        idx.append(("Grand Total", "Grand Total", ""))
    except Exception:
        pass
    display_cols = cols + ["Non RIS %", "RIS %"]
    df = pd.DataFrame(
        rows,
        columns=display_cols,
        index=pd.MultiIndex.from_tuples(idx, names=["State Cluster", "Brand", ""])
    )
    return df

def style_cluster_sheet_openpyxl(writer, sheet_name):
    ws = writer.sheets[sheet_name]
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_row = ws.max_row; max_col = ws.max_column
    for r in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in r:
            cell.border = border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    for row in range(2, max_row+1):
        brand_val = ws.cell(row=row, column=2).value
        try:
            if brand_val is None or str(brand_val).strip() == "":
                for col in range(1, max_col+1):
                    cell = ws.cell(row=row, column=col)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="B7CDE6", end_color="B7CDE6", fill_type="solid")
                a_cell = ws.cell(row=row, column=1); a_cell.alignment = Alignment(horizontal="left", indent=0)
            elif "Cluster Total" in str(brand_val):
                for col in range(1, max_col+1):
                    cell = ws.cell(row=row, column=col); cell.font = Font(bold=True); cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
            elif "Grand Total" in str(brand_val):
                for col in range(1, max_col+1):
                    cell = ws.cell(row=row, column=col); cell.font = Font(bold=True); cell.fill = PatternFill(start_color="9FB3C8", end_color="9FB3C8", fill_type="solid")
        except Exception:
            pass
    for row in range(2, max_row+1):
        bcell = ws.cell(row=row, column=2)
        if bcell.value and "Cluster Total" not in str(bcell.value) and "Grand Total" not in str(bcell.value):
            bcell.alignment = Alignment(horizontal="left", indent=2)
    for col in range(1, max_col+1):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 18

def apply_number_formats(ws, percent_cols_idx, number_cols_idx, header_row=1):
    max_row = ws.max_row
    for r in range(header_row + 1, max_row + 1):
        for c in percent_cols_idx:
            try:
                cell = ws.cell(row=r, column=c); cell.number_format = '0.00%'
            except Exception:
                pass
        for c in number_cols_idx:
            try:
                cell = ws.cell(row=r, column=c); cell.number_format = '#,##0'
            except Exception:
                pass


st.markdown('<div class="app-card">', unsafe_allow_html=True)
st.markdown('<span class="step-pill">STEP 1</span>', unsafe_allow_html=True)
st.markdown('<div class="section-title">Upload Files</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="section-subtitle">'
    'Upload the required Excel files to start the RIS analysis.'
    '</div>',
    unsafe_allow_html=True,
)

c1, c2, c3 = st.columns(3)

with c1:
    st.markdown("**Original.xlsx**")
    original_file = st.file_uploader("Upload Original.xlsx", type=["xlsx"], key="orig_up")

with c2:
    st.markdown("**FC Stat.xlsx**")
    fc_file = st.file_uploader("Upload FC Stat.xlsx", type=["xlsx"], key="fc_up")

with c3:
    st.markdown("**PM.xlsx (Purchase Master)**")
    pm_file = st.file_uploader("Upload PM.xlsx", type=["xlsx"], key="pm_up")

st.markdown("</div>", unsafe_allow_html=True)

if original_file and fc_file and pm_file:
    try:
        with st.spinner("Loading data..."):
            original_file.seek(0); Working = pd.read_excel(original_file)
            fc_file.seek(0); FC = pd.read_excel(fc_file)
            pm_file.seek(0)
            try:
                PM = pd.read_excel(pm_file, usecols="C,G", sheet_name=0)
            except Exception:
                PM = pd.read_excel(pm_file)

        st.success("Files loaded ✅")
        st.write(f"Original rows: {len(Working)} — FC rows: {len(FC)} — PM rows: {len(PM)}")

        if st.button("🔄 Process Data"):
            with st.spinner("Processing..."):
                # -- basic checks
                key_col_name = "fulfillment-center-id"
                if key_col_name not in Working.columns:
                    st.error(f"Missing column: {key_col_name}"); st.stop()
                if "quantity-shipped" not in Working.columns:
                    st.error("Missing 'quantity-shipped' column in Original file."); st.stop()

                # ensure numeric measure
                Working["quantity-shipped"] = pd.to_numeric(Working["quantity-shipped"], errors="coerce").fillna(0)

                # Merge FC info (unique keys only)
                FC_lookup = FC[["FC", "State", "Cluster"]].copy()
                FC_lookup.columns = ["lookup_key", "fulfillment_state", "cluster"]
                Working["_join"] = normalize(Working[key_col_name])
                FC_lookup["_join"] = normalize(FC_lookup["lookup_key"])
                FC_lookup = FC_lookup.drop_duplicates(subset=["_join"], keep="first")
                Working = Working.merge(
                    FC_lookup[["_join", "fulfillment_state", "cluster"]],
                    how="left",
                    left_on="_join",
                    right_on="_join"
                )
                Working.drop(columns=["_join"], inplace=True)

                # Normalize ship-state
                fc_states = FC["State"].dropna().astype(str).str.strip().tolist()
                canon_map = {normalize_text(s): s.strip() for s in fc_states if normalize_text(s) != ""}
                if "ship-state_original" not in Working.columns:
                    Working["ship-state_original"] = Working.get("ship-state", pd.NA)
                Working["ship-state"] = Working["ship-state_original"].apply(lambda x: safe_correct(x, canon_map))

                # --------------------------
                # Map DADRA & NAGAR HAVELI / DAMAN / DIU → Maharashtra
                # (applied to both ship-state and fulfillment_state)
                # --------------------------
                def map_to_maharashtra(x):
                    if pd.isna(x):
                        return x
                    s = str(x).strip().upper()
                    maha_aliases = {
                        "DADRA AND NAGAR HAVELI AND DAMAN AND DIU",
                        "DADRA AND NAGARA HAVELLI DAMAN AND DIU",
                        "DADRA & NAGAR HAVELI",
                        "DADRA AND NAGAR HAVELI",
                        "DADRA AND NAGARA HAVELLI",
                        "DAMAN",
                        "DIU",
                    }
                    if s in maha_aliases:
                        return "Maharashtra"
                    return x

                Working["ship-state"] = Working["ship-state"].apply(map_to_maharashtra)
                Working["fulfillment_state"] = Working["fulfillment_state"].apply(map_to_maharashtra)

                # Sheet2 VLOOKUP (A:C) using fulfillment-center-id if present
                try:
                    original_file.seek(0); xl = pd.ExcelFile(original_file)
                    sheet2 = xl.parse(xl.sheet_names[1]) if len(xl.sheet_names) > 1 else None
                except Exception:
                    sheet2 = None

                if sheet2 is not None and sheet2.shape[1] >= 3:
                    map_df = sheet2.iloc[:, [0, 2]].copy(); map_df.columns = ["key_raw", "state_cluster_raw"]
                    map_df["_k"] = map_df["key_raw"].astype(str).apply(lambda x: str(x).strip().upper())
                    mapping = dict(zip(map_df["_k"], map_df["state_cluster_raw"]))
                    Working["State Cluster"] = Working[key_col_name].astype(str).apply(
                        lambda x: mapping.get(str(x).strip().upper(), pd.NA)
                    )
                else:
                    Working["State Cluster"] = pd.NA

                # --------------------------
                # RIS status – use corrected ship-state & fulfillment_state
                # --------------------------
                def compute_ris_status(row):
                    ship_state = str(row.get("ship-state", "")).strip().replace(" ", "").upper()
                    fulfill_state = str(row.get("fulfillment_state", "")).strip().replace(" ", "").upper()
                    if ship_state and fulfill_state and ship_state == fulfill_state:
                        return "RIS"
                    return "Non RIS"

                Working["RIS Status"] = Working.apply(compute_ris_status, axis=1)

                # --------------------------
                # PM -> Brand mapping (Original.sku  <-> PM Column C -> Brand Column G)
                # --------------------------
                PM_lookup = pd.DataFrame()
                try:
                    if PM.shape[1] >= 2:
                        PM_lookup = PM.iloc[:, [0, 1]].copy()
                        PM_lookup.columns = ["lookup_key", "Brand_pm"]
                    else:
                        PM_lookup = pd.DataFrame()
                except Exception:
                    PM_lookup = pd.DataFrame()

                if not PM_lookup.empty:
                    PM_lookup["lookup_key_norm"] = PM_lookup["lookup_key"].apply(sku_normalize_raw)
                    PM_lookup["Brand_pm"] = PM_lookup["Brand_pm"].astype(str).str.strip()
                    PM_lookup = PM_lookup.drop_duplicates(subset=["lookup_key_norm"]).reset_index(drop=True)

                # Normalize Working.sku similarly
                if "sku" in Working.columns:
                    Working["_sku_norm"] = Working["sku"].apply(sku_normalize_raw)
                else:
                    Working["_sku_norm"] = ""

                # Ensure Brand present and preserve existing values: fill only missing
                if "Brand" not in Working.columns:
                    Working["Brand"] = ""
                Working["Brand"] = Working["Brand"].fillna("").astype(str)

                filled_count = 0
                if not PM_lookup.empty:
                    pm_map = dict(zip(PM_lookup["lookup_key_norm"], PM_lookup["Brand_pm"]))
                    pm_keys_set = set(pm_map.keys())
                    mask_unknown = Working["Brand"].str.strip().str.upper().isin(["", "NAN", "NONE", "UNKNOWN"])
                    match_mask = Working["_sku_norm"].isin(pm_keys_set) & mask_unknown
                    if match_mask.any():
                        for idx in Working[match_mask].index:
                            k = Working.at[idx, "_sku_norm"]
                            val = pm_map.get(k, "")
                            if val:
                                Working.at[idx, "Brand"] = val
                                filled_count += 1

                Working["Brand"] = Working["Brand"].replace("", "Unknown")

                st.markdown("### PM→Brand diagnostics")
                st.write("PM lookup rows (after normalize & dedup):", len(PM_lookup))
                st.write("Filled Brand rows from PM:", int(filled_count))
                st.write("Distinct Brands after fill:", int(Working["Brand"].nunique()))
                st.write(
                    "Rows remaining with Brand == 'Unknown':",
                    int((Working["Brand"].astype(str).str.strip().str.upper() == "UNKNOWN").sum())
                )

                if "_sku_norm" in Working.columns:
                    try:
                        Working.drop(columns=["_sku_norm"], inplace=True)
                    except Exception:
                        pass

                # --------------------------
                # Pivots
                # --------------------------
                detailed_pivot = pd.pivot_table(
                    Working,
                    values="quantity-shipped",
                    index=["Brand", "fulfillment_state", "ship-state"],
                    columns="RIS Status",
                    aggfunc="sum",
                    fill_value=0,
                )

                if not detailed_pivot.empty:
                    if "Non RIS" not in detailed_pivot.columns: detailed_pivot["Non RIS"] = 0
                    if "RIS" not in detailed_pivot.columns: detailed_pivot["RIS"] = 0
                    detailed_pivot["Grand Total"] = detailed_pivot["Non RIS"] + detailed_pivot["RIS"]

                    brand_totals = detailed_pivot.groupby(level=0).sum()
                    brand_totals.index = [(brand, f"{brand} Total", "") for brand in brand_totals.index]

                    state_totals = detailed_pivot.groupby(level=[0, 1]).sum()
                    state_totals.index = [(brand, f"{state} Total", "") for brand, state in state_totals.index]

                    detailed_pivot_with_totals = pd.concat(
                        [detailed_pivot, state_totals, brand_totals]
                    ).sort_index()

                    overall = detailed_pivot[["Non RIS", "RIS", "Grand Total"]].sum()
                    overall_idx = ("Grand Total", "", "")
                    overall_df = pd.DataFrame([overall.values], index=[overall_idx], columns=overall.index)
                    detailed_pivot_with_totals = pd.concat([detailed_pivot_with_totals, overall_df])
                else:
                    detailed_pivot_with_totals = detailed_pivot.copy()

                brand_summary = pd.pivot_table(
                    Working,
                    values="quantity-shipped",
                    index=["Brand"],
                    columns="RIS Status",
                    aggfunc="sum",
                    fill_value=0,
                    margins=False,
                )
                if "Non RIS" not in brand_summary.columns: brand_summary["Non RIS"] = 0
                if "RIS" not in brand_summary.columns: brand_summary["RIS"] = 0
                brand_summary["Grand Total"] = brand_summary["Non RIS"] + brand_summary["RIS"]
                brand_summary["Non RIS%"] = (brand_summary["Non RIS"] / brand_summary["Grand Total"]).fillna(0)
                brand_summary["RIS%"] = (brand_summary["RIS"] / brand_summary["Grand Total"]).fillna(0)

                overall_brand = brand_summary[["Non RIS", "RIS", "Grand Total"]].sum()
                overall_brand["Non RIS%"] = (
                    overall_brand["Non RIS"] / overall_brand["Grand Total"]
                    if overall_brand["Grand Total"] != 0 else 0
                )
                overall_brand["RIS%"] = (
                    overall_brand["RIS"] / overall_brand["Grand Total"]
                    if overall_brand["Grand Total"] != 0 else 0
                )
                overall_brand = overall_brand.rename("Grand Total")
                overall_brand_df = overall_brand.to_frame().T
                brand_summary = pd.concat([brand_summary, overall_brand_df])

                inventory_brand_summary = pd.pivot_table(
                    Working,
                    values="quantity-shipped",
                    index=["Brand"],
                    columns="RIS_by_Table",
                    aggfunc="sum",
                    fill_value=0,
                    margins=True,
                    margins_name="Grand Total",
                ) if "RIS_by_Table" in Working.columns else pd.DataFrame()

                if not inventory_brand_summary.empty:
                    if "Non RIS" not in inventory_brand_summary.columns: inventory_brand_summary["Non RIS"] = 0
                    if "RIS" not in inventory_brand_summary.columns: inventory_brand_summary["RIS"] = 0
                    inventory_brand_summary["Grand Total"] = inventory_brand_summary.get(
                        "Grand Total",
                        inventory_brand_summary["Non RIS"] + inventory_brand_summary["RIS"],
                    )
                    inventory_brand_summary["Non RIS%"] = (
                        inventory_brand_summary["Non RIS"] / inventory_brand_summary["Grand Total"]
                    ).fillna(0)
                    inventory_brand_summary["RIS%"] = (
                        inventory_brand_summary["RIS"] / inventory_brand_summary["Grand Total"]
                    ).fillna(0)

                cluster_pivot = pd.pivot_table(
                    Working,
                    values="quantity-shipped",
                    index=["State Cluster", "Brand"],
                    columns="RIS Status",
                    aggfunc="sum",
                    fill_value=0,
                )
                if not cluster_pivot.empty:
                    if "Non RIS" not in cluster_pivot.columns: cluster_pivot["Non RIS"] = 0
                    if "RIS" not in cluster_pivot.columns: cluster_pivot["RIS"] = 0
                    cluster_pivot["Grand Total"] = cluster_pivot["Non RIS"] + cluster_pivot["RIS"]

                display_cluster_df = build_cluster_display_with_pct(cluster_pivot)

            st.success("✅ Processing complete!")

            # --------------------------
            # UI: Tabs & displays
            # --------------------------
            tab1, tab2 = st.tabs(["📊 State-Based RIS Analysis", "🏭 Inventory Placement"])
            with tab1:
                st.subheader("Brand-Level Summary (with %)")
                st.dataframe(
                    brand_summary.style.format({
                        "Non RIS": "{:,.0f}",
                        "RIS": "{:,.0f}",
                        "Grand Total": "{:,.0f}",
                        "Non RIS%": "{:.2%}",
                        "RIS%": "{:.2%}",
                    }),
                    use_container_width=True,
                )

                st.subheader("Detailed (Brand → FC State → Ship State)")
                st.dataframe(
                    detailed_pivot_with_totals.style.format({
                        "Non RIS": "{:,.0f}",
                        "RIS": "{:,.0f}",
                        "Grand Total": "{:,.0f}",
                    }),
                    use_container_width=True,
                    height=400,
                )

                st.subheader("Cluster Analysis — Excel-like structure")
                if display_cluster_df.empty:
                    st.info("No cluster data available (Sheet2 missing or no matches).")
                else:
                    def style_rows(s):
                        label = s.name[1]
                        if label == "":
                            return ["background-color: #B7CDE6; font-weight: bold; text-align:left" for _ in s]
                        elif "Cluster Total" in str(label):
                            return ["background-color: #E0E0E0; font-weight: bold;" for _ in s]
                        elif "Grand Total" in str(label):
                            return ["background-color: #9FB3C8; font-weight: bold;" for _ in s]
                        else:
                            return ["" for _ in s]

                    st.dataframe(
                        display_cluster_df.style.apply(style_rows, axis=1).format({
                            "Non RIS": "{:,.0f}",
                            "RIS": "{:,.0f}",
                            "Grand Total": "{:,.0f}",
                            "Non RIS %": "{:.2%}",
                            "RIS %": "{:.2%}",
                        }),
                        use_container_width=True,
                        height=650,
                    )

            with tab2:
                st.header("Inventory Placement (Brand summary)")
                if not inventory_brand_summary.empty:
                    st.dataframe(
                        inventory_brand_summary.style.format({
                            "Non RIS": "{:,.0f}",
                            "RIS": "{:,.0f}",
                            "Grand Total": "{:,.0f}",
                            "Non RIS%": "{:.2%}",
                            "RIS%": "{:.2%}",
                        }),
                        use_container_width=True,
                    )
                else:
                    st.info("No Inventory Placement summary (RIS_by_Table missing).")

            # --------------------------
            # Downloads
            # --------------------------
            st.subheader("📥 Download Report (Excel)")
            col1, col2 = st.columns(2)
            with col1:
                out = BytesIO()
                with pd.ExcelWriter(out, engine="openpyxl") as writer:
                    brand_summary.to_excel(writer, sheet_name="Brand Summary")
                    detailed_pivot_with_totals.to_excel(writer, sheet_name="Detailed Analysis")
                    if not inventory_brand_summary.empty:
                        inventory_brand_summary.to_excel(writer, sheet_name="Inventory Brand Summary")
                    if not display_cluster_df.empty:
                        display_cluster_df.to_excel(writer, sheet_name="StateCluster Pivot")

                    for sheet_name in writer.sheets:
                        ws = writer.sheets[sheet_name]
                        for cell in ws[1]:
                            cell.font = Font(bold=True)

                        header_map = {
                            col_idx + 1: ws.cell(row=1, column=col_idx + 1).value
                            for col_idx in range(ws.max_column)
                        }

                        percent_cols = [
                            i for i, v in header_map.items()
                            if v in ["Non RIS%", "RIS%", "Non RIS %", "RIS %"]
                        ]

                        number_cols = [
                            i for i, v in header_map.items()
                            if isinstance(v, str)
                            and any(k in v for k in ["Non RIS", "RIS", "Grand Total"])
                            and "%" not in str(v)
                        ]

                        apply_number_formats(
                            ws,
                            percent_cols_idx=percent_cols,
                            number_cols_idx=number_cols,
                            header_row=1,
                        )

                        if sheet_name == "StateCluster Pivot":
                            style_cluster_sheet_openpyxl(writer, sheet_name)

                        for col in ws.columns:
                            max_len = 0
                            col_letter = get_column_letter(col[0].column)
                            for cell in col:
                                try:
                                    if cell.value:
                                        max_len = max(max_len, len(str(cell.value)))
                                except Exception:
                                    pass
                            ws.column_dimensions[col_letter].width = max(12, max_len + 2)

                out.seek(0)
                st.download_button(
                    "📦 Download Excel report",
                    data=out,
                    file_name="ris_cluster_report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            with col2:
                st.download_button(
                    "📄 Download Processed Data (CSV)",
                    data=Working.to_csv(index=False).encode("utf-8"),
                    file_name="processed_data.csv",
                    mime="text/csv",
                )

    except Exception as e:
        st.error("Error: " + str(e))
        st.exception(e)

else:
    st.info("Please upload **Original.xlsx**, **FC Stat.xlsx** and **PM.xlsx** above to proceed.")
    st.markdown(
        """
        - `Original.xlsx` must include **Sheet2** where Column A = `fulfillment-center-id`
          and Column C = **State Cluster** (used for VLOOKUP style).
        - `fulfillment-center-id` column is required in the Original sheet.
        - `quantity-shipped` is used as measure.
        """
    )
